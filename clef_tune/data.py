"""Load CSV/TSV/JSONL files into normalized records, and split train/eval.

Row numbers in messages are spreadsheet rows for CSV/TSV (the header is row 1) and line
numbers for JSONL.
"""

from __future__ import annotations

import csv
import datetime
import difflib
import io
import json
import random
from collections import defaultdict
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from clef_tune.config import Config, DataConfig
from clef_tune.labels import Label, LabelError, is_missing, parse_label
from clef_tune.schema import Question, Schema, SchemaError

TABLE_DELIMITERS = {".csv": ",", ".tsv": "\t"}
JSONL_SUFFIXES = {".jsonl", ".ndjson"}
EXCEL_SUFFIXES = {".xlsx"}
MAX_LISTED_ROWS = 20
URL_PREFIXES = ("http://", "https://")
CSV_FIELD_LIMIT = 2**31 - 1  # the largest value a C long holds on every platform

Row = tuple[int, dict[str, Any]]
Note = Callable[[str, str | None, str, int], None]
LabelNote = Callable[[str | None, str, int], None]  # a label that cannot be used


class DataError(ValueError):
    """A data file cannot be read at all."""


@dataclass
class Issue:
    level: Literal["error", "warning"]
    message: str
    rows: list[int] = field(default_factory=list)
    question: str | None = None

    def __str__(self) -> str:
        where = ""
        if self.rows:
            shown = ", ".join(str(row) for row in self.rows[:MAX_LISTED_ROWS])
            extra = len(self.rows) - MAX_LISTED_ROWS
            more = f" +{extra} more" if extra > 0 else ""
            where = f" (rows {shown}{more})"
        prefix = f"{self.question}: " if self.question else ""
        return f"{prefix}{self.message}{where}"


@dataclass
class Record:
    id: str
    state: Any
    questions: dict[str, Question]
    labels: dict[str, Label]
    images: list[Any] = field(default_factory=list)  # paths or URLs; Predictor rows: PIL too
    source_row: int = 0

    def to_json(self) -> dict[str, Any]:
        out: dict[str, Any] = {
            "id": self.id,
            "state": self.state,
            "questions": {qid: q.to_clef() for qid, q in self.questions.items()},
            "labels": {qid: label.to_json() for qid, label in self.labels.items()},
        }
        if self.images:
            out["images"] = list(self.images)
        return out


@dataclass
class LoadResult:
    records: list[Record]
    issues: list[Issue]
    columns: list[str]


def read_rows(path: Path) -> tuple[list[Row], list[str], list[Issue]]:
    """Read a CSV/TSV/JSONL file into numbered row dicts. Returns (rows, columns, issues)."""
    if not path.exists():
        raise DataError(f"{path}: file not found")
    suffix = path.suffix.lower()
    if suffix in TABLE_DELIMITERS:
        return _read_table(path, TABLE_DELIMITERS[suffix])
    if suffix in JSONL_SUFFIXES:
        return _read_jsonl(path)
    if suffix in EXCEL_SUFFIXES:
        return _read_excel(path)
    raise DataError(
        f"{path.name}: unsupported file type {suffix!r} (use .csv, .tsv, .xlsx or .jsonl)"
    )


def _decode(path: Path) -> tuple[str, list[Issue]]:
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8-sig"), []
    except UnicodeDecodeError:
        message = f"{path.name} is not UTF-8; read it as Windows-1252 (re-save as UTF-8 to be sure)"
        return raw.decode("cp1252", errors="replace"), [Issue("warning", message)]


def _read_table(path: Path, delimiter: str) -> tuple[list[Row], list[str], list[Issue]]:
    text, issues = _decode(path)
    csv.field_size_limit(CSV_FIELD_LIMIT)  # the 128 KB default rejects long documents
    records = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    rows, header, table_issues = _table_rows(records, path.name)
    return rows, header, issues + table_issues


def _table_rows(records: list[list[str]], name: str) -> tuple[list[Row], list[str], list[Issue]]:
    """Rows of a table whose first record is the header; numbers are spreadsheet row numbers."""
    issues: list[Issue] = []
    if not records or not any(cell.strip() for cell in records[0]):
        raise DataError(f"{name}: no header row")
    header = [cell.strip() for cell in records[0]]
    duplicates = sorted({column for column in header if header.count(column) > 1})
    if duplicates:
        raise DataError(f"{name}: duplicate column names {duplicates}")
    if "" in header:
        raise DataError(f"{name}: column {header.index('') + 1} has no name")
    rows: list[Row] = []
    bad_width: list[int] = []
    for number, values in enumerate(records[1:], start=2):
        if not any(value.strip() for value in values):
            continue  # blank line
        if len(values) != len(header):
            bad_width.append(number)
            continue
        rows.append((number, dict(zip(header, values, strict=True))))
    if bad_width:
        issues.append(Issue("error", f"rows must have {len(header)} fields", bad_width))
    return rows, header, issues


def _excel_cell(value: Any) -> str:
    """A cell as the text a CSV export would hold: whole numbers without ".0", dates as ISO."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime) and value.time() == datetime.time():
        return value.date().isoformat()
    if isinstance(value, datetime.date | datetime.time):
        return value.isoformat()
    return str(value)


def _read_excel(path: Path) -> tuple[list[Row], list[str], list[Issue]]:
    """The first sheet of an .xlsx workbook, read like a CSV: the first row is the header."""
    from openpyxl import load_workbook

    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # not a zip, or not a workbook
        raise DataError(f"{path.name}: not a readable Excel workbook ({exc})") from exc
    try:
        sheet = book.worksheets[0].iter_rows(values_only=True)
        cells = [[_excel_cell(value) for value in row] for row in sheet]
    finally:
        book.close()
    width = max((len(row) for row in cells), default=0)
    cells = [row + [""] * (width - len(row)) for row in cells]
    while width and all(not row[width - 1].strip() for row in cells):
        width -= 1  # Excel often reports formatted but empty columns at the right
    return _table_rows([row[:width] for row in cells], path.name)


def _read_jsonl(path: Path) -> tuple[list[Row], list[str], list[Issue]]:
    text, issues = _decode(path)
    rows: list[Row] = []
    columns: dict[str, None] = {}
    invalid: list[int] = []
    not_objects: list[int] = []
    # Split on "\n" only: str.splitlines() also breaks on U+2028, U+2029 and U+0085, which
    # json.dumps(ensure_ascii=False) writes raw inside strings.
    for number, raw_line in enumerate(text.split("\n"), start=1):
        line = raw_line.removesuffix("\r")
        if not line.strip():
            continue
        try:
            value = json.loads(line)
        except json.JSONDecodeError:
            invalid.append(number)
            continue
        if not isinstance(value, dict):
            not_objects.append(number)
            continue
        rows.append((number, value))
        columns.update(dict.fromkeys(value))
    if invalid:
        issues.append(Issue("error", "invalid JSON", invalid))
    if not_objects:
        issues.append(Issue("error", "each line must be a JSON object", not_objects))
    return rows, list(columns), issues


@dataclass
class _FlatLayout:
    state_columns: list[str]
    label_columns: dict[str, str]  # question id -> existing column


def _flat_layout(
    columns: list[str],
    cfg: DataConfig,
    schema: Schema,
    issues: list[Issue],
    for_prediction: bool = False,
) -> _FlatLayout:
    present = set(columns)

    def missing(column: str, what: str, question: str | None = None) -> None:
        hint = difflib.get_close_matches(column, columns, n=1)
        suffix = f" (did you mean '{hint[0]}'?)" if hint else ""
        issues.append(Issue("error", f"{what} '{column}' not found{suffix}", question=question))

    for qid in sorted(set(cfg.label_columns) - set(schema.questions)):
        issues.append(Issue("error", f"label_columns refers to unknown question '{qid}'"))
    label_columns: dict[str, str] = {}
    for qid in schema.questions:
        column = cfg.label_columns.get(qid, qid)
        if column in present:
            label_columns[qid] = column
        elif not for_prediction:
            missing(column, "label column", qid)
    if cfg.id_column and cfg.id_column not in present:
        missing(cfg.id_column, "id column")
    for column in cfg.image_columns:
        if column not in present:
            missing(column, "image column")
    if cfg.state_columns is not None:
        for column in cfg.state_columns:
            if column not in present:
                missing(column, "state column")
        state_columns = [c for c in cfg.state_columns if c in present]
    else:
        used = set(label_columns.values()) | set(cfg.image_columns) | {cfg.id_column}
        state_columns = [c for c in columns if c not in used]
    if not state_columns:
        issues.append(Issue("error", "no state columns: set data.state_columns"))
    return _FlatLayout(state_columns, label_columns)


def _resolve_image(value: str, data_dir: Path) -> str:
    if value.startswith(URL_PREFIXES):
        return value
    path = Path(value.replace("\\", "/"))
    return str(path if path.is_absolute() else (data_dir / path).resolve())


def _labels_for(
    number: int, pairs: Iterable[tuple[Question, Any]], note: Note, bad_label: LabelNote
) -> dict[str, Label]:
    labels: dict[str, Label] = {}
    for question, raw in pairs:
        try:
            label, warning = parse_label(question, raw)
        except LabelError as exc:
            bad_label(question.id, str(exc), number)
            continue
        if warning:
            note("warning", question.id, warning, number)
        if label is not None:
            labels[question.id] = label
    return labels


def _flat_record(
    number: int,
    row: dict[str, Any],
    cfg: DataConfig,
    schema: Schema,
    layout: _FlatLayout,
    data_dir: Path,
    note: Note,
    bad_label: LabelNote,
) -> Record | None:
    if cfg.id_column:
        raw_id = row.get(cfg.id_column)
        record_id = "" if is_missing(raw_id) else str(raw_id).strip()
        if not record_id:
            note("error", None, f"empty id in column '{cfg.id_column}'", number)
            return None
    else:
        record_id = f"row-{number}"
    values = {column: row.get(column) for column in layout.state_columns}
    state = values[layout.state_columns[0]] if len(layout.state_columns) == 1 else values
    pairs = ((schema.questions[qid], row.get(col)) for qid, col in layout.label_columns.items())
    labels = _labels_for(number, pairs, note, bad_label)
    images = [
        _resolve_image(str(row[column]).strip(), data_dir)
        for column in cfg.image_columns
        if column in row and not is_missing(row[column])
    ]
    return Record(record_id, state, schema.questions, labels, images, number)


def _native_record(
    number: int,
    row: dict[str, Any],
    schema: Schema | None,
    data_dir: Path,
    note: Note,
    question_cache: dict[str, dict[str, Question]],
    bad_label: LabelNote,
) -> Record | None:
    if "state" not in row:
        note("error", None, "native rows need a 'state'", number)
        return None
    if row.get("questions") is not None:
        key = json.dumps(row["questions"], sort_keys=True, default=str)
        if key not in question_cache:
            try:
                question_cache[key] = Schema.from_dict({"questions": row["questions"]}).questions
            except SchemaError as exc:
                note("error", None, f"invalid questions: {exc}", number)
                return None
        questions = question_cache[key]
    elif schema is not None:
        questions = schema.questions
    else:
        note("error", None, "row has no 'questions' and the config has no schema", number)
        return None
    pairs = []
    for qid, raw in (row.get("labels") or {}).items():
        question = questions.get(str(qid))
        if question is None:
            bad_label(str(qid), "label for a question the row does not ask", number)
            continue
        pairs.append((question, raw))
    raw_images = row.get("images") or []
    if isinstance(raw_images, str):
        raw_images = [raw_images]
    elif not isinstance(raw_images, list):
        note("error", None, "'images' must be a list of paths or URLs", number)
        return None
    labels = _labels_for(number, pairs, note, bad_label)
    images = [_resolve_image(str(value).strip(), data_dir) for value in raw_images]
    record_id = f"row-{number}" if is_missing(row.get("id")) else str(row["id"]).strip()
    return Record(record_id, row["state"], questions, labels, images, number)


def normalize_rows(
    rows: list[Row],
    columns: list[str],
    cfg: DataConfig,
    schema: Schema | None,
    data_dir: Path,
    for_prediction: bool = False,
    native_without_labels: bool = False,
) -> LoadResult:
    """Normalize rows into records. With ``for_prediction``, label columns are optional, rows
    without labels are kept, and labels that are present are still parsed, but one that cannot
    be used is a warning and is ignored. With ``native_without_labels`` (JSONL only), rows with
    a ``state`` and no ``labels`` are native."""
    issues: list[Issue] = []
    grouped: dict[tuple[str, str | None, str], list[int]] = defaultdict(list)

    def note(level: str, question: str | None, message: str, number: int) -> None:
        grouped[(level, question, message)].append(number)

    def bad_label(question: str | None, message: str, number: int) -> None:
        if for_prediction:  # predicting needs no labels: keep the row, drop the label
            note("warning", question, f"{message} (label ignored)", number)
        else:
            note("error", question, message, number)

    def is_native(row: dict[str, Any]) -> bool:
        if isinstance(row.get("labels"), dict):
            return True
        return native_without_labels and "state" in row and "labels" not in row

    layout = None
    if schema is not None and any(not is_native(r) for _, r in rows):
        layout = _flat_layout(columns, cfg, schema, issues, for_prediction)
    question_cache: dict[str, dict[str, Question]] = {}
    records: list[Record] = []
    unlabeled: list[int] = []
    for number, row in rows:
        if is_native(row):
            record = _native_record(number, row, schema, data_dir, note, question_cache, bad_label)
        elif schema is None or layout is None:
            message = "flat rows need a schema (add 'schema' or 'schema_file' to the config)"
            note("error", None, message, number)
            record = None
        else:
            record = _flat_record(number, row, cfg, schema, layout, data_dir, note, bad_label)
        if record is None:
            continue
        if not record.labels and not for_prediction:
            unlabeled.append(number)
            continue
        records.append(record)
    for (level, question, message), numbers in grouped.items():
        issues.append(Issue(level, message, numbers, question))  # type: ignore[arg-type]
    if unlabeled:
        message = f"{len(unlabeled)} rows have no labels and were dropped"
        issues.append(Issue("warning", message, unlabeled))
    return LoadResult(records, issues, columns)


def load_file(config: Config, path_value: str, *, for_prediction: bool = False) -> LoadResult:
    path = config.resolve_path(path_value)
    rows, columns, read_issues = read_rows(path)
    result = normalize_rows(
        rows,
        columns,
        config.data,
        config.get_schema(),
        path.parent,
        for_prediction,
        native_without_labels=for_prediction and path.suffix.lower() in JSONL_SUFFIXES,
    )
    result.issues = read_issues + result.issues
    return result


def _stratify_question(records: list[Record]) -> str | None:
    labeled = {qid for record in records for qid in record.labels}
    ordered = dict.fromkeys(qid for r in records for qid in [*r.questions, *r.labels])
    return next((qid for qid in ordered if qid in labeled), None)


def split_records(
    records: list[Record], eval_fraction: float, seed: int
) -> tuple[list[Record], list[Record]]:
    """Stratified, seeded split on the first labeled question; singleton classes stay in train."""
    key_question = _stratify_question(records)
    groups: dict[str | None, list[Record]] = defaultdict(list)
    for record in records:
        label = record.labels.get(key_question) if key_question else None
        groups[label.top() if label else None].append(record)
    rng = random.Random(seed)
    chosen: set[int] = set()
    for key in sorted(groups, key=lambda k: (k is None, str(k))):
        members = groups[key][:]
        rng.shuffle(members)
        count = round(len(members) * eval_fraction) if len(members) > 1 else 0
        chosen.update(id(member) for member in members[:count])
    if not chosen and len(records) > 1:
        largest = max(groups.values(), key=len)
        chosen.add(id(rng.choice(largest)))
    train = [record for record in records if id(record) not in chosen]
    evals = [record for record in records if id(record) in chosen]
    return train, evals
